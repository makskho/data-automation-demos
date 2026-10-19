"""Offline tests: parsing, pagination, checkpoint/resume, retries and the xlsx/csv writers.

The HTML fixtures were downloaded once from https://books.toscrape.com/ and are never fetched again.

Run:
    pytest -q
"""

import csv
import json
import os

import pytest
import requests
from openpyxl import load_workbook

import catalog_parser as cp

HERE = os.path.dirname(os.path.abspath(__file__))
FIXTURES = os.path.join(HERE, "fixtures")
BASE = "https://books.toscrape.com/"


def fixture(name):
    with open(os.path.join(FIXTURES, name), "r", encoding="utf-8") as stream:
        return stream.read()


PRODUCT_URL = BASE + "catalogue/a-light-in-the-attic_1000/index.html"


# ----------------------------------------------------------------------------- field rules


@pytest.mark.parametrize("text, expected", [
    ("£51.77", 51.77),
    ("£0.00", 0.0),
    ("Â£1,234.50", 1234.5),
    ("price: 10", 10.0),
    ("free", None),
    ("", None),
])
def test_parse_price(text, expected):
    assert cp.parse_price(text) == expected


@pytest.mark.parametrize("text, expected", [
    ("In stock (22 available)", 22),
    ("\n   In stock (1 available)\n", 1),
    ("Out of stock", 0),
    ("", 0),
])
def test_parse_availability(text, expected):
    assert cp.parse_availability(text) == expected


# ----------------------------------------------------------------------------- product page


def test_parse_product_full_record():
    record = cp.parse_product(fixture("product_a_light_in_the_attic.html"), PRODUCT_URL)
    assert record == {
        "title": "A Light in the Attic",
        "price": 51.77,
        "availability": 22,
        "rating": 3,
        "category": "Poetry",
        "upc": "a897fe39b1053632",
        "image_url": BASE + "media/cache/fe/72/fe72f0532301ec28892ae79a629a293c.jpg",
        "product_url": PRODUCT_URL,
    }
    assert isinstance(record["price"], float)
    assert isinstance(record["availability"], int)
    assert isinstance(record["rating"], int)


def test_parse_product_out_of_stock_and_missing_fields():
    html = fixture("product_a_light_in_the_attic.html")
    html = html.replace("In stock (22 available)", "Out of stock")
    html = html.replace('<p class="star-rating Three">', '<p class="star-rating">')
    html = html.replace("<th>UPC</th>", "<th>Code</th>")
    record = cp.parse_product(html, PRODUCT_URL)
    assert record["availability"] == 0
    assert record["rating"] is None
    assert record["upc"] is None
    assert record["title"] == "A Light in the Attic"  # the rest survives


def test_parse_product_rejects_a_foreign_page():
    with pytest.raises(cp.ParseError):
        cp.parse_product("<html><body><h1>404 Not Found</h1></body></html>", PRODUCT_URL)


# ----------------------------------------------------------------------------- listing and pagination


def test_listing_has_twenty_absolute_product_urls():
    urls = cp.parse_listing(fixture("catalogue_page_1.html"), BASE + "index.html")
    assert len(urls) == 20
    assert len(set(urls)) == 20
    assert urls[0] == PRODUCT_URL
    assert all(u.startswith(BASE + "catalogue/") and u.endswith("/index.html") for u in urls)


def test_next_page_found_and_resolved_from_site_root():
    # On the root page the link is relative to the root: catalogue/page-2.html
    assert cp.find_next_page(fixture("catalogue_page_1.html"), BASE + "index.html") == BASE + "catalogue/page-2.html"


def test_next_page_is_none_on_last_page():
    html = fixture("catalogue_page_last.html")
    assert "Page 50 of 50" in " ".join(html.split())
    assert cp.find_next_page(html, BASE + "catalogue/page-50.html") is None


def test_category_page_listing_resolves_relative_links():
    urls = cp.parse_listing(fixture("category_mystery.html"), BASE + "catalogue/category/books/mystery_3/index.html")
    assert len(urls) == 20
    assert all(u.startswith(BASE + "catalogue/") and "/category/" not in u for u in urls)


def test_categories_exclude_the_whole_catalogue_entry():
    categories = cp.parse_categories(fixture("catalogue_page_1.html"), BASE + "index.html")
    assert len(categories) == 50
    assert "Books" not in categories
    assert categories["Mystery"] == BASE + "catalogue/category/books/mystery_3/index.html"


# ----------------------------------------------------------------------------- crawl with a fake site


class FakeSession:
    """Serves the saved fixtures instead of the network and counts requests per URL."""

    def __init__(self):
        self.headers = {}
        self.calls = []

    def get(self, url, timeout=None):
        self.calls.append(url)
        if url == BASE:
            body = fixture("catalogue_page_1.html")
        elif url == BASE + "catalogue/page-2.html":
            body = fixture("catalogue_page_last.html")  # no 'next' link: the sandbox ends here
        elif url.endswith("/category/books/mystery_3/index.html"):
            body = fixture("category_mystery.html")
        elif url.endswith("/index.html") and "_" in url.rsplit("/", 2)[-2]:
            body = fixture("product_a_light_in_the_attic.html")
        else:
            return FakeResponse(404, b"")
        return FakeResponse(200, body.encode("utf-8"))

    def product_calls(self):
        return [u for u in self.calls if "/catalogue/" in u and "page-" not in u and "category" not in u]


class FakeResponse:
    def __init__(self, status, content, headers=None):
        self.status_code = status
        self.content = content
        self.headers = headers or {}


def make_fetcher(session):
    return cp.Fetcher(delay=0, session=session, sleep=lambda seconds: None)


def test_crawl_respects_limit_and_fetches_lazily():
    session = FakeSession()
    rows = cp.crawl(make_fetcher(session), BASE, limit=5)
    assert len(rows) == 5
    assert len(session.product_calls()) == 5
    assert BASE + "catalogue/page-2.html" not in session.calls  # one listing page was enough


def test_crawl_follows_pagination_to_the_end():
    session = FakeSession()
    rows = cp.crawl(make_fetcher(session), BASE)
    # page 1 and the last page both list 20 urls; fixtures overlap nowhere important, but urls are unique
    assert len(rows) == len({r["product_url"] for r in rows})
    assert BASE + "catalogue/page-2.html" in session.calls
    assert len(rows) > 20


def test_crawl_category_starts_from_the_category_page():
    session = FakeSession()
    rows = cp.crawl(make_fetcher(session), BASE, category="mystery", limit=3)
    assert len(rows) == 3
    assert any(u.endswith("/category/books/mystery_3/index.html") for u in session.calls)


def test_unknown_category_lists_the_available_ones():
    with pytest.raises(SystemExit) as stop:
        cp.crawl(make_fetcher(FakeSession()), BASE, category="Cooking with unicorns")
    assert "Mystery" in str(stop.value)


def test_checkpoint_resume_does_not_refetch(tmp_path):
    path = str(tmp_path / "run.checkpoint.jsonl")
    first = FakeSession()
    cp.crawl(make_fetcher(first), BASE, limit=3, checkpoint_path=path)
    assert len(first.product_calls()) == 3
    assert len(open(path, encoding="utf-8").read().splitlines()) == 3

    second = FakeSession()
    rows = cp.crawl(make_fetcher(second), BASE, limit=5, checkpoint_path=path)
    assert len(rows) == 5
    assert len(second.product_calls()) == 2  # only the two new ones
    assert len(open(path, encoding="utf-8").read().splitlines()) == 5


def test_checkpoint_survives_a_truncated_last_line(tmp_path):
    path = tmp_path / "run.checkpoint.jsonl"
    path.write_text(json.dumps({"product_url": "https://x/1", "title": "A"}) + '\n{"product_url": "https://x/2", "ti',
                    encoding="utf-8")
    assert list(cp.load_checkpoint(str(path))) == ["https://x/1"]


def test_failed_product_is_skipped_not_fatal():
    class Flaky(FakeSession):
        def get(self, url, timeout=None):
            if url.endswith("tipping-the-velvet_999/index.html"):
                return FakeResponse(404, b"")
            return super().get(url, timeout)

    rows = cp.crawl(make_fetcher(Flaky()), BASE, limit=4)
    assert len(rows) == 4
    assert all("tipping-the-velvet" not in r["product_url"] for r in rows)


# ----------------------------------------------------------------------------- retries


class ScriptedSession:
    def __init__(self, script):
        self.headers = {}
        self.script = list(script)
        self.calls = 0

    def get(self, url, timeout=None):
        self.calls += 1
        item = self.script.pop(0)
        if isinstance(item, Exception):
            raise item
        return item


def test_retries_with_exponential_backoff_then_succeeds():
    sleeps = []
    session = ScriptedSession([
        FakeResponse(503, b""),
        requests.ConnectionError("reset"),
        FakeResponse(429, b"", {"Retry-After": "7"}),
        FakeResponse(200, "£1".encode("utf-8")),
    ])
    fetcher = cp.Fetcher(delay=0, backoff=1.0, session=session, sleep=sleeps.append)
    assert fetcher.get(BASE) == "£1"
    assert session.calls == 4
    assert sleeps == [1.0, 2.0, 7.0]  # 1 s, 2 s, then the server's own Retry-After


def test_gives_up_after_the_configured_retries():
    session = ScriptedSession([FakeResponse(500, b"")] * 3)
    fetcher = cp.Fetcher(delay=0, retries=2, session=session, sleep=lambda s: None)
    with pytest.raises(requests.RequestException):
        fetcher.get(BASE)
    assert session.calls == 3


def test_404_is_not_retried():
    session = ScriptedSession([FakeResponse(404, b"")])
    fetcher = cp.Fetcher(delay=0, session=session, sleep=lambda s: None)
    with pytest.raises(requests.HTTPError):
        fetcher.get(BASE)
    assert session.calls == 1


def test_user_agent_is_set():
    session = FakeSession()
    cp.Fetcher(session=session, user_agent="my-agent/1.0")
    assert session.headers["User-Agent"] == "my-agent/1.0"


# ----------------------------------------------------------------------------- writers


ROWS = [
    {"title": "A", "price": 10.5, "availability": 3, "rating": 4, "category": "Poetry", "upc": "u1",
     "image_url": "https://x/a.jpg", "product_url": "https://x/a"},
    {"title": "B", "price": 20.0, "availability": 0, "rating": 1, "category": "Poetry", "upc": "u2",
     "image_url": "https://x/b.jpg", "product_url": "https://x/b"},
    {"title": "C", "price": 30.25, "availability": 8, "rating": 5, "category": "Travel", "upc": "u3",
     "image_url": "https://x/c.jpg", "product_url": "https://x/c"},
]


def test_summary_counts_and_average():
    stats = {s["category"]: s for s in cp.summarize(ROWS)}
    assert stats["Poetry"]["count"] == 2
    assert stats["Poetry"]["avg_price"] == 15.25
    assert stats["Poetry"]["min_price"] == 10.5 and stats["Poetry"]["max_price"] == 20.0
    assert stats["Travel"]["count"] == 1


def test_xlsx_types_and_formatting(tmp_path):
    path = str(tmp_path / "out" / "books.xlsx")
    cp.write_xlsx(ROWS, path)
    sheet = load_workbook(path)["Books"]

    assert [c.value for c in sheet[1]] == [title for _, title, _ in cp.COLUMNS]
    assert all(c.font.bold for c in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:H4"

    price, stock, rating = sheet["B2"], sheet["C2"], sheet["D2"]
    assert isinstance(price.value, float) and price.value == 10.5
    assert "0.00" in price.number_format
    assert isinstance(stock.value, int) and isinstance(rating.value, int)
    assert sheet["H2"].hyperlink is not None
    assert sheet.max_row == 4


def test_xlsx_summary_sheet(tmp_path):
    path = str(tmp_path / "books.xlsx")
    cp.write_xlsx(ROWS, path)
    summary = load_workbook(path)["Summary"]
    values = [[c.value for c in row] for row in summary.iter_rows()]
    assert values[0][:3] == ["Category", "Books", "Avg price, GBP"]
    assert values[1] == ["Poetry", 2, 15.25, 10.5, 20.0]
    assert values[-1][:3] == ["TOTAL", 3, 20.25]
    assert isinstance(values[1][2], float)


def test_xlsx_with_no_rows_still_opens(tmp_path):
    path = str(tmp_path / "empty.xlsx")
    cp.write_xlsx([], path)
    assert load_workbook(path)["Books"].max_row == 1


def test_csv_roundtrip_and_encoding(tmp_path):
    path = str(tmp_path / "books.csv")
    rows = [dict(ROWS[0], title="Война и мир, том 1")]
    cp.write_csv(rows, path)
    with open(path, "rb") as stream:
        assert stream.read(3) == b"\xef\xbb\xbf"  # BOM for Excel
    with open(path, "r", encoding="utf-8-sig", newline="") as stream:
        back = list(csv.DictReader(stream))
    assert back[0]["title"] == "Война и мир, том 1"
    assert list(back[0]) == cp.FIELDS
    assert not os.path.exists(path + ".tmp")


def test_real_fixture_roundtrip_to_xlsx(tmp_path):
    record = cp.parse_product(fixture("product_a_light_in_the_attic.html"), PRODUCT_URL)
    path = str(tmp_path / "one.xlsx")
    cp.write_xlsx([record], path)
    sheet = load_workbook(path)["Books"]
    assert sheet["A2"].value == "A Light in the Attic"
    assert sheet["B2"].value == 51.77
    assert sheet["F2"].value == "a897fe39b1053632"


# ----------------------------------------------------------------------------- CLI


def test_cli_rejects_bad_limit():
    assert cp.main(["--limit", "0"]) == 2


def test_cli_end_to_end_with_fake_site(tmp_path, monkeypatch):
    session = FakeSession()
    original = cp.Fetcher

    def patched(**kwargs):
        kwargs["session"] = session
        kwargs["sleep"] = lambda s: None
        kwargs["delay"] = 0
        return original(**kwargs)

    monkeypatch.setattr(cp, "Fetcher", patched)
    out = str(tmp_path / "res" / "books.xlsx")
    assert cp.main(["--limit", "4", "--out", out]) == 0
    assert load_workbook(out)["Books"].max_row == 5
    assert os.path.exists(out.replace(".xlsx", ".csv"))
    assert os.path.exists(out.replace(".xlsx", ".checkpoint.jsonl"))


def test_limit_equal_to_page_size_does_not_fetch_the_next_listing_page():
    session = FakeSession()
    rows = cp.crawl(make_fetcher(session), BASE, limit=20)
    assert len(rows) == 20
    assert BASE + "catalogue/page-2.html" not in session.calls
