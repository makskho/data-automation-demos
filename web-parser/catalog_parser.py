"""Catalogue parser: crawl a paginated product catalogue and write it to Excel and CSV.

Target of this demo: https://books.toscrape.com/ (a public sandbox built for scraping practice).
Everything that knows the markup of that site lives in the four ``parse_*`` / ``find_next_page``
functions. The rest (polite fetching, retries, checkpoint, writers, CLI) is site-independent.

Run:
    python catalog_parser.py --limit 60 --out sample-output/books.xlsx
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import random
import re
import sys
import time
from typing import Callable, Dict, Iterable, Iterator, List, Optional
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/"
USER_AGENT = "catalog-parser-demo/1.0 (portfolio sample; polite crawler, 1 request per ~0.7 s)"

# (record key, column title, width)
COLUMNS = [
    ("title", "Title", 48),
    ("price", "Price, GBP", 12),
    ("availability", "In stock, pcs", 14),
    ("rating", "Rating (1-5)", 13),
    ("category", "Category", 20),
    ("upc", "UPC", 20),
    ("image_url", "Image URL", 60),
    ("product_url", "Product URL", 60),
]
FIELDS = [key for key, _, _ in COLUMNS]

RATING_WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5}


class ParseError(ValueError):
    """The page does not look like the page this parser was written for."""


# --------------------------------------------------------------------------- parsing


def soup_of(html: str) -> BeautifulSoup:
    return BeautifulSoup(html, "lxml")


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def parse_price(text: str) -> Optional[float]:
    """'£51.77' -> 51.77. Returns None when there is no number at all."""
    match = re.search(r"\d[\d\s,]*(?:\.\d+)?", text or "")
    if not match:
        return None
    return float(match.group(0).replace(",", "").replace(" ", ""))


def parse_availability(text: str) -> int:
    """'In stock (22 available)' -> 22; 'Out of stock' or no number -> 0."""
    match = re.search(r"\((\d+)\s+available\)", text or "")
    return int(match.group(1)) if match else 0


def parse_rating(tag) -> Optional[int]:
    """<p class="star-rating Three"> -> 3."""
    if tag is None:
        return None
    for css_class in tag.get("class", []):
        if css_class.lower() in RATING_WORDS:
            return RATING_WORDS[css_class.lower()]
    return None


def parse_listing(html: str, page_url: str) -> List[str]:
    """Absolute product URLs found on one catalogue (or category) page, in page order."""
    urls = []
    for pod in soup_of(html).select("article.product_pod h3 a[href]"):
        urls.append(urljoin(page_url, pod["href"]))
    return urls


def find_next_page(html: str, page_url: str) -> Optional[str]:
    """Absolute URL of the next page, or None on the last page."""
    link = soup_of(html).select_one("ul.pager li.next a[href]")
    return urljoin(page_url, link["href"]) if link else None


def parse_categories(html: str, page_url: str) -> Dict[str, str]:
    """Sidebar categories as {name: absolute url}; the top 'Books' entry is the whole catalogue."""
    result = {}
    for link in soup_of(html).select("div.side_categories ul ul a[href]"):
        result[clean_text(link.get_text())] = urljoin(page_url, link["href"])
    return result


def parse_product(html: str, url: str) -> Dict[str, object]:
    """One product page -> one flat record. Missing optional fields become None."""
    soup = soup_of(html)
    heading = soup.select_one("div.product_main h1")
    if heading is None:
        raise ParseError(f"no product title on {url}")

    info = {}
    for row in soup.select("table.table tr"):
        th, td = row.find("th"), row.find("td")
        if th is not None and td is not None:
            info[clean_text(th.get_text())] = clean_text(td.get_text())

    price_tag = soup.select_one("div.product_main p.price_color")
    price = parse_price(price_tag.get_text()) if price_tag else parse_price(info.get("Price (incl. tax)", ""))

    availability_tag = soup.select_one("div.product_main p.availability")
    availability_text = availability_tag.get_text() if availability_tag else info.get("Availability", "")

    crumbs = soup.select("ul.breadcrumb li")
    category = None
    if len(crumbs) >= 3:
        category = clean_text(crumbs[-2].get_text())

    image = soup.select_one("div.item img[src]") or soup.select_one("#product_gallery img[src]")

    return {
        "title": clean_text(heading.get_text()),
        "price": price,
        "availability": parse_availability(availability_text),
        "rating": parse_rating(soup.select_one("div.product_main p.star-rating")),
        "category": category,
        "upc": info.get("UPC"),
        "image_url": urljoin(url, image["src"]) if image else None,
        "product_url": url,
    }


# --------------------------------------------------------------------------- polite fetching


class Fetcher:
    """GET with a fixed delay, jitter, retries and exponential backoff.

    Retries on connection errors, timeouts, 429 and 5xx. Honours ``Retry-After``.
    A 404 or any other 4xx is not retried: it is a fact about the page, not a hiccup.
    """

    RETRY_STATUSES = {429, 500, 502, 503, 504}

    def __init__(
        self,
        delay: float = 0.7,
        retries: int = 4,
        backoff: float = 1.5,
        timeout: float = 20.0,
        user_agent: str = USER_AGENT,
        session: Optional[requests.Session] = None,
        sleep: Callable[[float], None] = time.sleep,
        log: Callable[[str], None] = lambda message: None,
    ):
        self.delay = delay
        self.retries = retries
        self.backoff = backoff
        self.timeout = timeout
        self.session = session or requests.Session()
        self.session.headers["User-Agent"] = user_agent
        self.sleep = sleep
        self.log = log
        self._last_request = 0.0
        self.requests_made = 0

    def _pause(self) -> None:
        if self.delay <= 0 or not self._last_request:
            return
        wait = self.delay + random.uniform(0, self.delay * 0.3) - (time.monotonic() - self._last_request)
        if wait > 0:
            self.sleep(wait)

    def get(self, url: str) -> str:
        last_problem = "unknown"
        for attempt in range(self.retries + 1):
            self._pause()
            self._last_request = time.monotonic()
            self.requests_made += 1
            retry_after = None
            try:
                response = self.session.get(url, timeout=self.timeout)
            except (requests.ConnectionError, requests.Timeout) as error:
                last_problem = f"{type(error).__name__}"
            else:
                if response.status_code == 200:
                    # The site sends no charset; requests would guess latin-1 and break the pound sign.
                    return response.content.decode("utf-8", errors="replace")
                last_problem = f"HTTP {response.status_code}"
                if response.status_code not in self.RETRY_STATUSES:
                    raise requests.HTTPError(f"{last_problem} for {url}")
                header = response.headers.get("Retry-After", "")
                retry_after = float(header) if header.isdigit() else None
            if attempt < self.retries:
                pause = retry_after if retry_after is not None else self.backoff * (2 ** attempt)
                self.log(f"  {last_problem} on {url}; retry {attempt + 1}/{self.retries} in {pause:.1f}s")
                self.sleep(pause)
        raise requests.RequestException(f"giving up on {url} after {self.retries + 1} tries ({last_problem})")


# --------------------------------------------------------------------------- crawling


def iter_listing_urls(fetcher: Fetcher, start_url: str) -> Iterator[str]:
    """Walk the pagination from ``start_url`` and yield product URLs lazily (no page is fetched
    before it is needed, so a small --limit costs only a few requests)."""
    seen_pages = set()
    page_url: Optional[str] = start_url
    while page_url and page_url not in seen_pages:
        seen_pages.add(page_url)
        html = fetcher.get(page_url)
        yield from parse_listing(html, page_url)
        page_url = find_next_page(html, page_url)


def resolve_start_url(fetcher: Fetcher, base_url: str, category: Optional[str]) -> str:
    if not category:
        return base_url
    categories = parse_categories(fetcher.get(base_url), base_url)
    for name, url in categories.items():
        if name.lower() == category.strip().lower():
            return url
    raise SystemExit(f"Unknown category {category!r}. Available: {', '.join(sorted(categories))}")


def load_checkpoint(path: str) -> Dict[str, Dict[str, object]]:
    done: Dict[str, Dict[str, object]] = {}
    if not os.path.exists(path):
        return done
    with open(path, "r", encoding="utf-8") as stream:
        for line in stream:
            line = line.strip()
            if not line:
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                continue  # a half-written last line after an interrupted run
            if isinstance(record, dict) and record.get("product_url"):
                done[record["product_url"]] = record
    return done


def crawl(
    fetcher: Fetcher,
    base_url: str = BASE_URL,
    category: Optional[str] = None,
    limit: Optional[int] = None,
    checkpoint_path: Optional[str] = None,
    log: Callable[[str], None] = lambda message: None,
) -> List[Dict[str, object]]:
    """Return product records in catalogue order. Records already in the checkpoint are reused."""
    done = load_checkpoint(checkpoint_path) if checkpoint_path else {}
    if done:
        log(f"checkpoint: {len(done)} products already collected")
    start_url = resolve_start_url(fetcher, base_url, category)
    rows: List[Dict[str, object]] = []
    failed: List[str] = []
    seen = set()

    checkpoint = open(checkpoint_path, "a", encoding="utf-8") if checkpoint_path else None
    try:
        for url in iter_listing_urls(fetcher, start_url):
            if url in seen:
                continue
            seen.add(url)
            if url in done:
                rows.append(done[url])
                if limit is not None and len(rows) >= limit:
                    break  # break before the generator fetches one more listing page
                continue
            try:
                record = parse_product(fetcher.get(url), url)
            except (requests.RequestException, ParseError) as error:
                failed.append(url)
                log(f"  skipped {url}: {error}")
                continue
            rows.append(record)
            if checkpoint:
                checkpoint.write(json.dumps(record, ensure_ascii=False) + "\n")
                checkpoint.flush()
            if len(rows) % 10 == 0:
                log(f"  {len(rows)} products")
            if limit is not None and len(rows) >= limit:
                break
    finally:
        if checkpoint:
            checkpoint.close()
    if failed:
        log(f"{len(failed)} product pages failed; run again to retry them")
    return rows


# --------------------------------------------------------------------------- output


def summarize(rows: Iterable[Dict[str, object]]) -> List[Dict[str, object]]:
    """Per category: count, average / min / max price. Sorted by count, then name."""
    counts: Dict[str, int] = {}
    groups: Dict[str, List[float]] = {}
    for row in rows:
        key = row.get("category") or "(none)"
        counts[key] = counts.get(key, 0) + 1
        groups.setdefault(key, [])
        if isinstance(row.get("price"), (int, float)):
            groups[key].append(float(row["price"]))
    summary = []
    for name in sorted(groups, key=lambda n: (-counts[n], n)):
        prices = groups[name]
        summary.append({
            "category": name,
            "count": counts[name],
            "avg_price": round(sum(prices) / len(prices), 2) if prices else None,
            "min_price": min(prices) if prices else None,
            "max_price": max(prices) if prices else None,
        })
    return summary


HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
THIN = Side(style="thin", color="BFBFBF")
PRICE_FORMAT = '"£"#,##0.00'


def _style_header(sheet, titles: List[str], widths: List[int]) -> None:
    for index, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 22
    sheet.freeze_panes = "A2"


def build_workbook(rows: List[Dict[str, object]]) -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = "Books"
    _style_header(sheet, [title for _, title, _ in COLUMNS], [width for _, _, width in COLUMNS])
    for row in rows:
        sheet.append([row.get(key) for key in FIELDS])

    column = {key: index + 1 for index, key in enumerate(FIELDS)}
    last_row = len(rows) + 1
    for number in range(2, last_row + 1):
        sheet.cell(row=number, column=column["price"]).number_format = PRICE_FORMAT
        sheet.cell(row=number, column=column["availability"]).number_format = "0"
        sheet.cell(row=number, column=column["rating"]).alignment = Alignment(horizontal="center")
        link = sheet.cell(row=number, column=column["product_url"])
        if link.value:
            link.hyperlink = link.value
            link.font = Font(color="0563C1", underline="single")
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}{max(last_row, 2)}"

    summary = book.create_sheet("Summary")
    _style_header(summary, ["Category", "Books", "Avg price, GBP", "Min price, GBP", "Max price, GBP"],
                  [24, 10, 16, 16, 16])
    stats = summarize(rows)
    for item in stats:
        summary.append([item["category"], item["count"], item["avg_price"], item["min_price"], item["max_price"]])
    prices = [float(r["price"]) for r in rows if isinstance(r.get("price"), (int, float))]
    total = [
        "TOTAL", len(rows),
        round(sum(prices) / len(prices), 2) if prices else None,
        min(prices) if prices else None,
        max(prices) if prices else None,
    ]
    summary.append(total)
    for number in range(2, summary.max_row + 1):
        for col in (3, 4, 5):
            summary.cell(row=number, column=col).number_format = PRICE_FORMAT
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
        cell.border = Border(top=THIN)
    return book


def _atomic(path: str, writer: Callable[[str], None]) -> None:
    folder = os.path.dirname(os.path.abspath(path))
    os.makedirs(folder, exist_ok=True)
    temp = path + ".tmp"
    writer(temp)
    os.replace(temp, path)


def write_xlsx(rows: List[Dict[str, object]], path: str) -> None:
    _atomic(path, lambda temp: build_workbook(rows).save(temp))


def write_csv(rows: List[Dict[str, object]], path: str) -> None:
    def writer(temp: str) -> None:
        # utf-8-sig so that Excel opens the file with the right encoding.
        with open(temp, "w", encoding="utf-8-sig", newline="") as stream:
            out = csv.DictWriter(stream, fieldnames=FIELDS, extrasaction="ignore")
            out.writeheader()
            out.writerows(rows)
    _atomic(path, writer)


# --------------------------------------------------------------------------- CLI


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Crawl a product catalogue into Excel and CSV.")
    parser.add_argument("--base-url", default=BASE_URL, help="catalogue start page (default: %(default)s)")
    parser.add_argument("--limit", type=int, default=None, help="stop after N products (default: all)")
    parser.add_argument("--category", default=None, help="only this category, e.g. Mystery")
    parser.add_argument("--out", default="output/books.xlsx", help="xlsx path; the CSV is written next to it")
    parser.add_argument("--delay", type=float, default=0.7, help="seconds between requests (default: %(default)s)")
    parser.add_argument("--retries", type=int, default=4, help="retries per request (default: %(default)s)")
    parser.add_argument("--timeout", type=float, default=20.0, help="request timeout, seconds")
    parser.add_argument("--checkpoint", default=None, help="checkpoint file (default: <out>.checkpoint.jsonl)")
    parser.add_argument("--fresh", action="store_true", help="ignore and delete an existing checkpoint")
    parser.add_argument("--user-agent", default=USER_AGENT)
    return parser


def main(argv: Optional[List[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    if args.limit is not None and args.limit < 1:
        print("--limit must be at least 1", file=sys.stderr)
        return 2
    stem, _ = os.path.splitext(args.out)
    checkpoint = args.checkpoint or stem + ".checkpoint.jsonl"
    if args.fresh and os.path.exists(checkpoint):
        os.remove(checkpoint)
    os.makedirs(os.path.dirname(os.path.abspath(checkpoint)), exist_ok=True)

    def log(message: str) -> None:
        print(message, file=sys.stderr, flush=True)

    fetcher = Fetcher(delay=args.delay, retries=args.retries, timeout=args.timeout,
                      user_agent=args.user_agent, log=log)
    started = time.monotonic()
    try:
        rows = crawl(fetcher, args.base_url, args.category, args.limit, checkpoint, log)
    except KeyboardInterrupt:
        log("interrupted; collected products are in the checkpoint, run the same command to continue")
        return 130
    except requests.RequestException as error:
        log(f"stopped: {error}; run the same command to continue from the checkpoint")
        return 1
    if not rows:
        log("nothing collected; no files written")
        return 1
    write_xlsx(rows, stem + ".xlsx")
    write_csv(rows, stem + ".csv")
    log(f"done: {len(rows)} products, {fetcher.requests_made} requests, "
        f"{time.monotonic() - started:.0f}s -> {stem}.xlsx / {stem}.csv")
    return 0


if __name__ == "__main__":
    sys.exit(main())
