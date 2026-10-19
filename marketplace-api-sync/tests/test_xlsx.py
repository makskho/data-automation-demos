from datetime import datetime

from openpyxl import load_workbook

from ozon_sync.sheets import (
    PRODUCT_COLUMNS,
    PRODUCTS_SHEET,
    POSTINGS_SHEET,
    read_edit_sheet,
    write_postings,
    write_products,
)
from ozon_sync.cli import day_window
from datetime import date


def test_products_xlsx_has_typed_cells_and_formatting(tmp_path, client):
    path = write_products(tmp_path / "p.xlsx", client.fetch_catalog())
    wb = load_workbook(path)
    ws = wb[PRODUCTS_SHEET]
    assert [c.value for c in ws[1]] == [c.header for c in PRODUCT_COLUMNS]
    assert ws.max_row == 15  # header + 14 products
    row = {c.key: ws.cell(row=2, column=i).value for i, c in enumerate(PRODUCT_COLUMNS, start=1)}
    assert row["offer_id"] == "KR-350-WH" and isinstance(row["offer_id"], str)
    assert isinstance(row["product_id"], int)
    assert isinstance(row["price"], (int, float)) and row["price"] == 390  # a number, not the text "390"
    assert isinstance(row["fbs_present"], int) and row["fbs_present"] == 120
    assert row["name"] == "Кружка керамическая 350 мл, белая"
    # formatting a seller sees
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref == "A1:J15"
    assert ws["A1"].font.bold and ws["D2"].fill.fgColor.rgb.endswith("FFF2CC")  # editable column tinted
    assert ws["D2"].number_format == "#,##0.00" and ws["H2"].number_format == "0"
    assert ws["A2"].number_format == "@"  # offer_id stays text in Excel
    assert ws.data_validations.dataValidation  # stock column rejects negatives in Excel itself
    assert "Как редактировать" in wb.sheetnames


def test_products_roundtrip_through_reader(tmp_path, client):
    path = write_products(tmp_path / "p.xlsx", client.fetch_catalog())
    rows = read_edit_sheet(path)
    assert len(rows) == 14 and not any(r.parse_errors for r in rows)
    first = rows[0]
    assert (first.row_no, first.offer_id, int(first.price), first.fbs_present) == (2, "KR-350-WH", 390, 120)


def test_postings_xlsx_types_and_local_time(tmp_path, client):
    since, to = day_window(date(2026, 10, 1), date(2026, 10, 2), 3)
    postings = client.list_postings(since, to)
    path = write_postings(tmp_path / "o.xlsx", postings, utc_offset_hours=3)
    ws = load_workbook(path)[POSTINGS_SHEET]
    headers = [c.value for c in ws[1]]
    assert "Принят в обработку" in headers
    rows = [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 3  # posting 1 has one line, posting 2 has two
    first = rows[0]
    assert isinstance(first["Принят в обработку"], datetime)
    assert first["Принят в обработку"] == datetime(2026, 10, 1, 9, 12, 44)  # 06:12:44 UTC shown as Moscow time
    assert first["Статус"] == "Доставлен"
    assert isinstance(first["Кол-во"], int) and first["Кол-во"] == 2
    assert first["Цена за шт."] == 390 and first["Сумма"] == 780  # numbers, line total = qty * price
    assert ws["D2"].number_format == "DD.MM.YYYY HH:MM"
    assert ws.freeze_panes == "B2"
