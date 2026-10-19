"""FBS postings: date filter, pagination by offset, status filter, CLI export."""
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook

from ozon_sync.cli import day_window, main
from ozon_sync.client import format_ozon_time, parse_ozon_time


def numbers(postings):
    return [p.posting_number for p in postings]


def test_window_is_inclusive_of_both_days_in_seller_timezone():
    since, to = day_window(date(2026, 10, 8), date(2026, 10, 8), utc_offset_hours=3)
    assert format_ozon_time(since) == "2026-10-07T21:00:00.000Z"
    assert format_ozon_time(to) == "2026-10-08T20:59:59.999Z"


def test_date_filter_keeps_only_postings_inside_the_range(client):
    since, to = day_window(date(2026, 10, 5), date(2026, 10, 8), 3)
    postings = client.list_postings(since, to)
    # Oct 3 and Oct 9+ are outside; 2026-10-07T21:59:59Z is 00:59 on Oct 8 in Moscow, so it is inside
    assert numbers(postings) == ["48230981-0001-1", "48257404-0001-1", "48257404-0002-1", "48290215-0001-1"]
    assert all(since <= p.in_process_at <= to for p in postings)


def test_timezone_decides_which_day_a_posting_belongs_to(client):
    # 2026-10-07T21:59:59Z: still Oct 7 in UTC, already Oct 8 in Moscow
    in_utc = numbers(client.list_postings(*day_window(date(2026, 10, 8), date(2026, 10, 8), 0)))
    in_msk = numbers(client.list_postings(*day_window(date(2026, 10, 8), date(2026, 10, 8), 3)))
    assert "48257404-0002-1" not in in_utc and "48257404-0002-1" in in_msk
    assert in_utc == ["48290215-0001-1"]


def test_boundaries_are_inclusive(client):
    exact = parse_ozon_time("2026-10-08T00:00:00.000Z")
    postings = client.list_postings(exact, exact)
    assert numbers(postings) == ["48290215-0001-1"]


def test_empty_range_returns_nothing(client):
    assert client.list_postings(*day_window(date(2026, 11, 1), date(2026, 11, 30), 3)) == []


def test_status_filter(client):
    since, to = day_window(date(2026, 10, 1), date(2026, 10, 31), 3)
    cancelled = client.list_postings(since, to, status="cancelled")
    assert numbers(cancelled) == ["48201377-0001-1"]
    assert len(client.list_postings(since, to)) == 12


def test_pagination_by_offset_while_has_next(fake, client):
    since, to = day_window(date(2026, 10, 1), date(2026, 10, 31), 3)
    postings = client.list_postings(since, to, page=5)
    assert len(postings) == 12 and len(set(numbers(postings))) == 12
    calls = fake.calls_to("/v3/posting/fbs/list")
    assert [c["offset"] for c in calls] == [0, 5, 10]  # 5 + 5 + 2, no extra request after has_next=false
    assert calls[0]["filter"]["since"] == "2026-09-30T21:00:00.000Z"  # 2026-10-01 00:00 Moscow
    assert calls[0]["filter"]["to"] == "2026-10-31T20:59:59.999Z"


def test_lines_and_fields_are_parsed(client):
    since, to = day_window(date(2026, 10, 2), date(2026, 10, 2), 3)
    (posting,) = client.list_postings(since, to)
    assert posting.status == "delivered" and posting.city == "Казань"
    assert posting.in_process_at == datetime(2026, 10, 2, 9, 40, 10, tzinfo=timezone.utc)
    assert [(l.offer_id, l.quantity, int(l.price)) for l in posting.lines] == [("PL-150-GR", 1, 1290), ("PD-5050-2", 1, 560)]


def test_cli_exports_the_range_to_xlsx(fake, tmp_path, capsys):
    out = tmp_path / "orders.xlsx"
    code = main(["export-postings", "--from", "2026-10-01", "--to", "2026-10-07", "-o", str(out)], transport=fake)
    assert code == 0
    assert "Отправлений: 5, строк товаров: 7" in capsys.readouterr().out
    ws = load_workbook(out)["Отправления"]
    assert ws.max_row == 8  # header + 7 product lines
    numbers_in_file = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert "48257404-0002-1" not in numbers_in_file  # it is Oct 8 in Moscow


def test_cli_rejects_reversed_or_malformed_dates(fake, tmp_path, capsys):
    assert main(["export-postings", "--from", "2026-10-09", "--to", "2026-10-01", "-o", str(tmp_path / "x.xlsx")], transport=fake) == 2
    assert fake.calls == []
    with pytest.raises(SystemExit):
        main(["export-postings", "--from", "01.10.2026", "--to", "2026-10-02"], transport=fake)
