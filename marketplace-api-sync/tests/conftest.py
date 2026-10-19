"""Shared helpers. All data is synthetic; nothing here touches the network."""
import pytest
from openpyxl import load_workbook

from ozon_sync.client import OzonClient
from ozon_sync.fake import FakeOzon
from ozon_sync.sheets import PRODUCT_COLUMNS, PRODUCTS_SHEET, write_products

COL = {c.key: i for i, c in enumerate(PRODUCT_COLUMNS, start=1)}
WRITE_PATHS = ("/v2/products/stocks", "/v1/product/import/prices")


def synthetic_products(n: int) -> list[dict]:
    """n generated catalogue entries in the fixture shape (for batching / pagination volume)."""
    return [
        {
            "product_id": 5_000_000 + i,
            "sku": 9_000_000 + i,
            "offer_id": f"GEN-{i:05d}",
            "name": f"Тестовый товар {i}",
            "archived": False,
            "price": {"price": "100", "old_price": "0", "min_price": "0", "net_price": "50", "currency_code": "RUB"},
            "stocks": [{"type": "fbs", "present": 10, "reserved": 0}, {"type": "fbo", "present": 0, "reserved": 0}],
        }
        for i in range(n)
    ]


@pytest.fixture
def fake():
    return FakeOzon()


@pytest.fixture
def client(fake):
    return OzonClient(fake)


@pytest.fixture
def exported(tmp_path, client):
    """The seller's workbook, exactly as export-products writes it."""
    path = tmp_path / "products.xlsx"
    write_products(path, client.fetch_catalog())
    return path


def edit_cells(path, edits: dict[str, dict[str, object]], extra: list[dict] = ()):
    """Set cells by (offer_id, field) in the exported workbook; `extra` appends new rows."""
    wb = load_workbook(path)
    ws = wb[PRODUCTS_SHEET]
    rows = {ws.cell(row=r, column=COL["offer_id"]).value: r for r in range(2, ws.max_row + 1)}
    for offer_id, fields in edits.items():
        for key, value in fields.items():
            ws.cell(row=rows[offer_id], column=COL[key]).value = value
    for item in extra:
        row = ws.max_row + 1
        for key, value in item.items():
            ws.cell(row=row, column=COL[key]).value = value
    wb.save(path)
    return rows


def write_calls(fake: FakeOzon):
    return [(p, payload) for p, payload in fake.calls if p in WRITE_PATHS]
