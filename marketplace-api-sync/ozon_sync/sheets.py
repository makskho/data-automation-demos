"""Excel input/output: formatted exports and reading the sheet the seller edits."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .errors import SheetError
from .models import Posting, Product, to_decimal

PRODUCTS_SHEET = "Товары"
POSTINGS_SHEET = "Отправления"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EDIT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_RU = {
    "awaiting_packaging": "Ожидает сборки",
    "awaiting_deliver": "Ожидает отгрузки",
    "delivering": "Доставляется",
    "delivered": "Доставлен",
    "cancelled": "Отменён",
}


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    width: int
    fmt: str = "General"
    editable: bool = False


PRODUCT_COLUMNS = [
    Column("offer_id", "Артикул (offer_id)", 16, "@"),
    Column("product_id", "ID товара", 14, "0"),
    Column("name", "Название", 46),
    Column("price", "Цена", 12, "#,##0.00", True),
    Column("old_price", "Цена до скидки", 15, "#,##0.00", True),
    Column("min_price", "Минимальная цена", 17, "#,##0.00", True),
    Column("currency", "Валюта", 9),
    Column("fbs_present", "Остаток FBS", 13, "0", True),
    Column("fbs_reserved", "Резерв FBS", 12, "0"),
    Column("fbo_present", "Остаток FBO", 13, "0"),
]

POSTING_COLUMNS = [
    Column("posting_number", "Номер отправления", 20, "@"),
    Column("order_number", "Номер заказа", 16, "@"),
    Column("status", "Статус", 20),
    Column("in_process_at", "Принят в обработку", 19, "DD.MM.YYYY HH:MM"),
    Column("shipment_date", "Отгрузить до", 19, "DD.MM.YYYY HH:MM"),
    Column("delivery_method", "Способ доставки", 24),
    Column("city", "Город", 18),
    Column("offer_id", "Артикул", 14, "@"),
    Column("name", "Товар", 44),
    Column("quantity", "Кол-во", 8, "0"),
    Column("price", "Цена за шт.", 13, "#,##0.00"),
    Column("amount", "Сумма", 13, "#,##0.00"),
    Column("currency", "Валюта", 9),
]

EDITABLE_KEYS = ("price", "old_price", "min_price", "fbs_present")


def _number(value: Decimal) -> int | float:
    return int(value) if value == value.to_integral_value() else float(value)


def _style_sheet(ws, columns: Sequence[Column], row_count: int) -> None:
    for index, column in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = column.width
        for row in range(2, row_count + 2):
            body = ws.cell(row=row, column=index)
            body.number_format = column.fmt
            body.border = BORDER
            if column.editable:
                body.fill = EDIT_FILL
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(row_count + 1, 2)}"


def write_products(path: str | Path, products: Sequence[Product]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = PRODUCTS_SHEET
    ws.append([c.header for c in PRODUCT_COLUMNS])
    for p in products:
        ws.append(
            [
                p.offer_id,
                p.product_id,
                p.name,
                _number(p.price),
                _number(p.old_price),
                _number(p.min_price),
                p.currency,
                p.fbs_present,
                p.fbs_reserved,
                p.fbo_present,
            ]
        )
    _style_sheet(ws, PRODUCT_COLUMNS, len(products))
    if products:
        stock_col = get_column_letter(1 + [c.key for c in PRODUCT_COLUMNS].index("fbs_present"))
        rule = DataValidation(
            type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
            errorTitle="Остаток", error="Остаток — целое число не меньше 0",
        )
        ws.add_data_validation(rule)
        rule.add(f"{stock_col}2:{stock_col}{len(products) + 1}")
    guide = wb.create_sheet("Как редактировать")
    for line in (
        "Меняйте только жёлтые столбцы: Цена, Цена до скидки, Минимальная цена, Остаток FBS.",
        "Пустая ячейка = не менять. Цена до скидки = 0 — убрать зачёркнутую цену.",
        "Остальные столбцы справочные, их значения при загрузке игнорируются.",
        "Артикул (offer_id) не меняйте и не удаляйте: по нему строка сопоставляется с товаром.",
        "Сначала запустите команду update без --apply: она покажет таблицу изменений и ничего не отправит.",
    ):
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    return _save(wb, path)


def write_postings(path: str | Path, postings: Sequence[Posting], utc_offset_hours: float = 3) -> Path:
    """One row per product line. Times are shown in the seller's timezone (default Moscow, UTC+3)."""
    tz = timezone(timedelta(hours=utc_offset_hours))

    def local(moment: datetime | None):
        return moment.astimezone(tz).replace(tzinfo=None) if moment else None

    wb = Workbook()
    ws = wb.active
    ws.title = POSTINGS_SHEET
    ws.append([c.header for c in POSTING_COLUMNS])
    rows = 0
    for posting in postings:
        for line in posting.lines:
            ws.append(
                [
                    posting.posting_number,
                    posting.order_number,
                    STATUS_RU.get(posting.status, posting.status),
                    local(posting.in_process_at),
                    local(posting.shipment_date),
                    posting.delivery_method,
                    posting.city,
                    line.offer_id,
                    line.name,
                    line.quantity,
                    _number(line.price),
                    _number(line.price * line.quantity),
                    line.currency,
                ]
            )
            rows += 1
    _style_sheet(ws, POSTING_COLUMNS, rows)
    return _save(wb, path)


def _save(wb: Workbook, path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target


# ---- reading the seller's sheet ---------------------------------------------------------------


@dataclass
class SheetRow:
    row_no: int
    offer_id: str
    price: Decimal | None = None
    old_price: Decimal | None = None
    min_price: Decimal | None = None
    fbs_present: int | None = None
    parse_errors: list[tuple[str, str]] = field(default_factory=list)  # (code, text)


def _norm(text: object) -> str:
    return " ".join(str(text).split()).casefold()


def _header_map() -> dict[str, str]:
    mapping: dict[str, str] = {}
    for column in PRODUCT_COLUMNS:
        mapping[_norm(column.header)] = column.key
        mapping[_norm(column.key)] = column.key
    mapping[_norm("offer_id")] = "offer_id"
    mapping[_norm("Артикул")] = "offer_id"
    mapping[_norm("Остаток")] = "fbs_present"
    return mapping


def _offer_id(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def read_edit_sheet(path: str | Path) -> list[SheetRow]:
    """Read the editable columns. Type problems are collected per row, never raised."""
    try:
        wb = load_workbook(path, data_only=True)
    except FileNotFoundError:
        raise SheetError(f"файл не найден: {path}") from None
    except Exception as exc:  # corrupt / not an xlsx
        raise SheetError(f"не удалось открыть {path} как .xlsx: {exc}") from None
    ws = wb[PRODUCTS_SHEET] if PRODUCTS_SHEET in wb.sheetnames else wb.worksheets[0]
    header_to_key = _header_map()
    columns: dict[str, int] = {}
    for index, cell in enumerate(ws[1]):
        key = header_to_key.get(_norm(cell.value)) if cell.value is not None else None
        if key and key not in columns:
            columns[key] = index
    if "offer_id" not in columns:
        raise SheetError("в первой строке не найден столбец «Артикул (offer_id)»")

    rows: list[SheetRow] = []
    for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(_blank(v) for v in values):
            continue

        def cell(key: str):
            idx = columns.get(key)
            return values[idx] if idx is not None and idx < len(values) else None

        row = SheetRow(row_no=row_no, offer_id=_offer_id(cell("offer_id")))
        for key in ("price", "old_price", "min_price"):
            raw = cell(key)
            if _blank(raw):
                continue
            try:
                setattr(row, key, to_decimal(raw))
            except ValueError:
                row.parse_errors.append(("bad_number", f"{key}: «{raw}» не число"))
        raw = cell("fbs_present")
        if not _blank(raw):
            try:
                number = to_decimal(raw)
                if number != number.to_integral_value():
                    row.parse_errors.append(("fractional_stock", f"fbs_present: {raw} — остаток должен быть целым"))
                else:
                    row.fbs_present = int(number)
            except ValueError:
                row.parse_errors.append(("bad_number", f"fbs_present: «{raw}» не число"))
        rows.append(row)
    return rows
