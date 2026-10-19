"""Make two seller-style edits of sample-output/products.xlsx for the README walkthrough.

    python tools/make_demo_edits.py
"""
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent / "sample-output"
SOURCE = ROOT / "products.xlsx"
HEADERS = {"offer_id": 1, "price": 4, "old_price": 5, "min_price": 6, "fbs_present": 8}


def rows_by_offer(ws):
    return {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}


def edit(target: Path, changes: dict[str, dict[str, object]], extra_rows: list[dict[str, object]] = ()):
    wb = load_workbook(SOURCE)
    ws = wb["Товары"]
    index = rows_by_offer(ws)
    for offer_id, fields in changes.items():
        for key, value in fields.items():
            ws.cell(row=index[offer_id], column=HEADERS[key]).value = value
    for item in extra_rows:
        row = ws.max_row + 1
        for key, value in item.items():
            ws.cell(row=row, column=HEADERS[key]).value = value
    wb.save(target)
    print(target)


def main():
    edit(
        ROOT / "products_edited.xlsx",
        {
            "KR-350-WH": {"price": 420, "fbs_present": 150},
            "PL-150-GR": {"price": 1390, "old_price": 1890},
            "SH-LED-3M": {"fbs_present": 40},
            "TR-USB-C1": {"fbs_present": 190},
        },
    )
    edit(
        ROOT / "products_edited_with_errors.xlsx",
        {
            "KR-350-WH": {"price": 420},
            "CT-BAM-S": {"fbs_present": -5},
            "OR-BOX-3": {"price": "тысяча"},
            "LM-DSK-W": {"price": 12900},
        },
        [{"offer_id": "NO-SUCH-SKU", "price": 500}],
    )


if __name__ == "__main__":
    main()
