"""Сценарии диалога через настоящий Dispatcher; Bot API подменён, сеть не используется."""
from datetime import datetime, timedelta, timezone
from io import BytesIO

from aiogram.methods import AnswerCallbackQuery, EditMessageText, SendDocument, SendMessage
from openpyxl import load_workbook

from conftest import ADMIN_ID, USER_ID


async def fill_form(h, name="Иван Петров", phone="8 (999) 123-45-67", task="Нужен Telegram-бот для записи клиентов", **who):
    await h.press("lead", **who)
    await h.say(name, **who)
    await h.say(phone, **who)
    await h.say(task, **who)


async def test_start_shows_welcome_and_menu(harness):
    await harness.say("/start")
    [msg] = harness.session.sent(SendMessage, USER_ID)
    assert msg.text == harness.faq.welcome
    buttons = [b.callback_data for row in msg.reply_markup.inline_keyboard for b in row]
    assert buttons == ["n:services", "n:prices", "n:payment", "n:contacts", "lead"]


async def test_faq_navigation_forward_and_back(harness):
    await harness.press("n:services")
    await harness.press("n:services-bots")
    await harness.press("n:services-bots-faq")
    edits = harness.session.sent(EditMessageText)
    assert edits[0].text.startswith("Что мы делаем")
    assert edits[1].text.startswith("Боты с меню")
    assert "Меню с вопросами" in edits[2].text
    back = [b.callback_data for row in edits[2].reply_markup.inline_keyboard for b in row]
    assert back == ["lead", "n:services-bots", "n:root"]

    await harness.press("n:root")
    assert harness.session.sent(EditMessageText)[-1].text == harness.faq.welcome
    assert len(harness.session.sent(AnswerCallbackQuery)) == 4


async def test_unknown_node_returns_to_menu(harness):
    await harness.press("n:deleted-from-file")
    assert harness.session.sent(EditMessageText)[-1].text == harness.faq.welcome


async def test_full_form_saves_lead_and_notifies_admin(harness):
    await fill_form(harness, username="ivan_p")

    [lead] = harness.repo.all()
    assert (lead.name, lead.phone, lead.user_id, lead.username) == ("Иван Петров", "+79991234567", USER_ID, "ivan_p")
    assert lead.task == "Нужен Telegram-бот для записи клиентов"

    [notice] = harness.session.sent(SendMessage, ADMIN_ID)
    for part in ("№1", "Иван Петров", "+79991234567", "Нужен Telegram-бот", "@ivan_p"):
        assert part in notice.text
    assert "Заявка №1 принята" in harness.session.texts(USER_ID)[-1]


async def test_invalid_input_keeps_user_on_the_same_step(harness):
    await harness.press("lead")
    await harness.say("Иван123")  # имя с цифрами
    await harness.say("Иван")
    await harness.say("не телефон")
    await harness.say("123")
    await harness.say("+79991234567")
    await harness.say("коротко")
    assert harness.repo.count() == 0
    texts = harness.session.texts(USER_ID)
    assert "без цифр" in texts[1]
    assert "Укажите телефон" in texts[2]
    assert "Не удалось распознать номер" in texts[3] and "Не удалось распознать номер" in texts[4]
    assert "Коротко опишите задачу" in texts[5]
    assert "минимум 10" in texts[6]

    await harness.say("Нужен сайт-визитка для кофейни")  # форма продолжает работать после ошибок
    assert harness.repo.count() == 1
    assert harness.repo.all()[0].name == "Иван"


async def test_non_text_answer_is_refused(harness):
    await harness.press("lead")
    await harness.say_photo()
    assert "ответьте текстом" in harness.session.texts(USER_ID)[-1]


async def test_cancel_resets_the_form(harness):
    await harness.press("lead")
    await harness.say("Иван")
    await harness.say("/cancel")
    assert "отменена" in harness.session.texts(USER_ID)[-1]
    await harness.say("+79991234567")  # без формы это обычный текст
    assert "/start" in harness.session.texts(USER_ID)[-1]
    assert harness.repo.count() == 0


async def test_start_in_the_middle_of_a_form_resets_it(harness):
    await harness.press("lead")
    await harness.say("Иван")
    await harness.say("/start")
    await harness.say("+79991234567")
    assert harness.repo.count() == 0


async def test_command_is_not_taken_as_form_answer(harness):
    await harness.press("lead")
    await harness.say("/export")  # не администратор
    assert "только администратору" in harness.session.texts(USER_ID)[-1]
    assert harness.repo.count() == 0


async def test_antiflood_blocks_second_form(harness):
    await fill_form(harness)
    sent_before = len(harness.session.texts(USER_ID))
    await harness.press("lead")
    reply = harness.session.texts(USER_ID)[sent_before:]
    assert len(reply) == 1 and "через 10 мин" in reply[0]
    await harness.say("Пётр")  # форма не открылась
    assert harness.repo.count() == 1


async def test_antiflood_allows_after_the_interval(harness):
    old = datetime.now(timezone.utc) - timedelta(minutes=11)
    harness.repo.add(USER_ID, "client", "Иван", "+79991234567", "Прошлая заявка давно", created_at=old)
    await fill_form(harness, name="Пётр")
    assert harness.repo.count() == 2


async def test_antiflood_is_per_user(harness):
    await fill_form(harness)
    await fill_form(harness, name="Анна", user_id=2002)
    assert harness.repo.count() == 2


async def test_antiflood_rechecked_at_submit(harness):
    await harness.press("lead")
    await harness.say("Иван")
    await harness.say("+79991234567")
    # заявка ушла из другого диалога, пока форма была открыта
    harness.repo.add(USER_ID, "client", "Иван", "+79991234567", "Параллельная заявка", created_at=datetime.now(timezone.utc))
    await harness.say("Описание задачи для проверки гонки")
    assert harness.repo.count() == 1
    assert "через" in harness.session.texts(USER_ID)[-1]
    assert harness.session.sent(SendMessage, ADMIN_ID) == []


async def test_lead_is_kept_when_admin_notification_fails(harness):
    harness.session.fail_chat_ids.add(ADMIN_ID)
    await fill_form(harness)
    assert harness.repo.count() == 1
    assert "принята" in harness.session.texts(USER_ID)[-1]


async def test_export_for_admin_sends_xlsx(harness):
    await fill_form(harness)
    await harness.say("/export", user_id=ADMIN_ID, username="boss")
    [doc] = harness.session.sent(SendDocument, ADMIN_ID)
    assert doc.document.filename == "leads.xlsx"
    ws = load_workbook(BytesIO(doc.document.data)).active
    assert ws.max_row == 2 and ws["C2"].value == "Иван Петров"


async def test_export_without_leads(harness):
    await harness.say("/export", user_id=ADMIN_ID)
    assert "пока нет" in harness.session.texts(ADMIN_ID)[-1]
    assert harness.session.sent(SendDocument) == []


async def test_export_denied_for_regular_user(harness):
    await fill_form(harness)
    await harness.say("/export")
    assert harness.session.sent(SendDocument) == []
    assert "только администратору" in harness.session.texts(USER_ID)[-1]


async def test_plain_text_outside_form_points_to_start(harness):
    await harness.say("привет")
    assert "/start" in harness.session.texts(USER_ID)[-1]
