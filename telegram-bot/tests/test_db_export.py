from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook

from bot.db import LeadRepo
from bot.export import HEADERS, build_xlsx

T0 = datetime(2026, 5, 1, 12, 0, tzinfo=timezone.utc)


def test_add_and_read_back(repo):
    lead = repo.add(10, "ivan", "Иван", "+79991234567", "Нужен бот", created_at=T0)
    assert lead.id == 1
    assert repo.count() == 1
    assert repo.all() == [lead]
    assert lead.created_at == T0


def test_data_survives_reopening(tmp_path):
    path = str(tmp_path / "x.db")
    LeadRepo(path).add(1, None, "Анна", "+79990000000", "Описание задачи", created_at=T0)
    again = LeadRepo(path)
    assert again.count() == 1
    assert again.all()[0].username is None


def test_sql_is_parametrized(repo):
    nasty = "x'); DROP TABLE leads;--"
    repo.add(1, "u", nasty, "+79990000000", nasty, created_at=T0)
    assert repo.all()[0].name == nasty


def test_flood_wait_logic(repo):
    assert repo.flood_wait(1, 10, now=T0) is None  # заявок ещё не было
    repo.add(1, "u", "Иван", "+79990000000", "Описание задачи", created_at=T0)
    assert repo.flood_wait(1, 10, now=T0 + timedelta(minutes=4)) == timedelta(minutes=6)
    assert repo.flood_wait(1, 10, now=T0 + timedelta(minutes=10)) is None
    assert repo.flood_wait(2, 10, now=T0 + timedelta(minutes=1)) is None  # другой пользователь
    assert repo.flood_wait(1, 0, now=T0 + timedelta(seconds=1)) is None  # защита выключена


def test_export_xlsx_contents(repo):
    repo.add(10, "ivan", "Иван", "+79991234567", "Первая задача", created_at=T0)
    repo.add(11, None, "Анна", "+79990000001", "Вторая задача\nв две строки", created_at=T0 + timedelta(hours=1))
    data = build_xlsx(repo.all())
    ws = load_workbook(BytesIO(data)).active
    assert ws.title == "Заявки"
    assert tuple(c.value for c in ws[1]) == HEADERS
    assert ws.max_row == 3
    assert [c.value for c in ws[2]] == [1, "2026-05-01 12:00", "Иван", "+79991234567", "Первая задача", "@ivan", 10]
    assert ws[3][5].value in ("", None)
    assert ws[3][4].value == "Вторая задача\nв две строки"
    assert ws.freeze_panes == "A2"


def test_export_empty_has_only_header():
    ws = load_workbook(BytesIO(build_xlsx([]))).active
    assert ws.max_row == 1


def test_formula_like_text_stays_text(repo):
    repo.add(1, None, "Иван", "+79990000000", "=HYPERLINK(\"http://x\",\"y\")", created_at=T0)
    cell = load_workbook(BytesIO(build_xlsx(repo.all()))).active["E2"]
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")
