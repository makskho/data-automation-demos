"""Выгрузка заявок в .xlsx."""
from __future__ import annotations

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .db import Lead

HEADERS = ("№", "Дата (UTC)", "Имя", "Телефон", "Задача", "Username", "Telegram ID")
_WIDTHS = (6, 18, 22, 18, 60, 20, 14)


def _text(cell, value: str) -> None:
    """Пишет значение как текст: строка, начинающаяся с '=', не станет формулой."""
    cell.value = value
    cell.data_type = "s"


def build_xlsx(leads: Iterable[Lead]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for lead in leads:
        ws.append(
            [
                lead.id,
                lead.created_at.strftime("%Y-%m-%d %H:%M"),
                lead.name,
                lead.phone,
                lead.task,
                f"@{lead.username}" if lead.username else "",
                lead.user_id,
            ]
        )
        row = ws[ws.max_row]
        for cell in row:
            if isinstance(cell.value, str):
                _text(cell, cell.value)
        row[4].alignment = Alignment(wrap_text=True, vertical="top")
    for idx, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
