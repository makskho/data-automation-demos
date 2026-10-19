#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Build the translation / editing portfolio samples (DOCX) from the text sources in src/.

Usage:   python build_samples.py          (needs python-docx; see README for the local .venv)

Sources (src/*.md)
  01_ru-en_samovar.md         RU->EN translation pair, glossary, notes
  02_en-ru_groundwater.md     EN->RU translation pair, glossary, notes
  03_literary-editing_ru.md   author's draft + edits in one marked-up text, editor's note
  04_business-editing_en.md   same for an English business document, change summary

Markup used in the editing sources (one source of truth: the draft and the edited text are both
derived from it, so they can never drift apart):
  {-deleted-}   text present in the draft, removed by the editor
  {+inserted+}  text added by the editor;  {+¶+} = the editor splits the paragraph here
  ^[comment]    a margin comment anchored to the previous piece of text
Everything else is unchanged text.

Output goes to docx/.  The script exits non-zero if an automated QA check fails.
"""
import re
import sys
import zipfile
from collections import Counter
from datetime import datetime
from itertools import count
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "src"
OUT = ROOT / "docx"
AUTHOR = "Максим"
INITIALS = "М"
BUILD_ISO = "2026-10-19T12:00:00Z"
BUILD_DT = datetime(2026, 10, 19, 12, 0, 0)

RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x7A, 0x2E)
GREY = RGBColor(0x59, 0x59, 0x59)
HEAD = RGBColor(0x1F, 0x38, 0x64)

FOOTER = {
    "ru": "Демонстрационный образец для портфолио · автор: Максим · не является работой для клиента",
    "en": "Portfolio demonstration sample · author: Maxim · not client work",
}

errors = []


def fail(msg):
    errors.append(msg)
    print("  QA FAIL:", msg)


# --------------------------------------------------------------------------- parsing
def split_sections(text):
    """-> (header dict, {section name: [lines]})"""
    header, sections, cur = {}, {}, None
    for line in text.splitlines():
        if line.startswith("## "):
            cur = line[3:].strip()
            sections[cur] = []
        elif cur is None:
            if ":" in line:
                k, v = line.split(":", 1)
                header[k.strip()] = v.strip()
        else:
            sections[cur].append(line)
    return header, sections


def parse_table(lines):
    rows = []
    for ln in lines:
        ln = ln.strip()
        if not ln.startswith("|"):
            continue
        cells = [c.strip() for c in ln.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows[0], rows[1:]


def parse_note(lines):
    """-> [(kind, text)] ; kinds: h, h2, p, li"""
    out = []
    for ln in lines:
        ln = ln.strip()
        if not ln:
            continue
        m = re.match(r"@(h2|h|p|li)\s+(.*)", ln)
        out.append((m.group(1), m.group(2)) if m else ("p", ln))
    return out


def parse_translation(path):
    header, sec = split_sections(path.read_text(encoding="utf-8"))
    src_lang, tgt_lang = header["direction"].split("-")
    pairs, cur = [], None
    for ln in sec["Text"]:
        if ln.startswith("### "):
            cur = {}
            pairs.append(cur)
        elif cur is not None:
            m = re.match(r"(RU|EN):\s*(.*)", ln)
            if m:
                cur[m.group(1).lower()] = m.group(2).strip()
    gl_head, gl_rows = parse_table(sec["Glossary"])
    notes = [ln.strip()[2:] for ln in sec.get("Notes", []) if ln.strip().startswith("- ")]
    return dict(header=header, src=src_lang, tgt=tgt_lang, pairs=pairs,
                glossary=gl_rows, gl_head=gl_head, notes=notes)


TOKEN = re.compile(r"\{-(.*?)-\}|\{\+(.*?)\+\}|\^\[(.*?)\]", re.S)


def tokenize(text):
    out, pos = [], 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            out.append(("keep", text[pos:m.start()]))
        if m.group(1) is not None:
            out.append(("del", m.group(1)))
        elif m.group(2) is not None:
            out.append(("ins", m.group(2)))
        else:
            out.append(("com", m.group(3)))
        pos = m.end()
    if pos < len(text):
        out.append(("keep", text[pos:]))
    return out


def parse_edit(path):
    header, sec = split_sections(path.read_text(encoding="utf-8"))
    paras = []
    for block in re.split(r"\n\s*\n", "\n".join(sec["Text"]).strip()):
        block = " ".join(block.split("\n")).strip()
        if not block:
            continue
        m = re.match(r"@(t|h1|h2|li)\s+(.*)", block, re.S)
        kind, body = (m.group(1), m.group(2)) if m else ("body", block)
        paras.append((kind, tokenize(body)))
    note = parse_note(sec.get("Note", []))
    summary = parse_table(sec["Summary"]) if "Summary" in sec and sec["Summary"] and any(
        l.strip().startswith("|") for l in sec["Summary"]) else None
    return dict(header=header, paras=paras, note=note, summary=summary)


def derive(tokens, which):
    """which = 'draft' | 'edited' -> text, with '\n' marking inserted paragraph breaks."""
    s = []
    for t, x in tokens:
        if t == "keep":
            s.append(x)
        elif t == "del" and which == "draft":
            s.append(x)
        elif t == "ins":
            if x == "¶":
                s.append(" " if which == "draft" else "\n")
            elif which == "edited":
                s.append(x)
    return "".join(s)


# --------------------------------------------------------------------------- QA helpers
ROMAN = {"I": 1, "V": 5, "X": 10, "L": 50}


def roman_to_int(s):
    total = 0
    for i, ch in enumerate(s):
        v = ROMAN[ch]
        total += -v if i + 1 < len(s) and ROMAN[s[i + 1]] > v else v
    return total


def numbers(text, lang):
    t = text.replace("`", "")
    if lang == "ru":
        t = re.sub(r"\b([IVXL]{1,6})(?=\s+(?:век|столет))", lambda m: str(roman_to_int(m.group(1))), t)
        pat = r"\d+(?:[   ]\d{3})*(?:,\d+)?"
    else:
        t = re.sub(r"(\d+)(?:st|nd|rd|th)\b", r"\1", t)
        pat = r"\d+(?:,\d{3})*(?:\.\d+)?"
    toks = []
    for m in re.findall(pat, t):
        v = re.sub(r"[   ]", "", m)
        v = v.replace(",", ".") if lang == "ru" else v.replace(",", "")
        toks.append(v.lstrip("0") or "0")
    return Counter(toks)


def stems(phrase):
    return [w[:max(3, len(w) - 2)].lower() for w in re.findall(r"[\w'’]+", phrase)]


def glossary_used(target_text, translation_cell):
    low = target_text.lower()
    for alt in translation_cell.split("/"):
        st = stems(alt)
        if st and all(s in low for s in st):
            return True
    return False


def qa_translation(tr):
    """Returns a dict of QA findings; records failures."""
    src, tgt = tr["src"], tr["tgt"]
    res = dict(paras=len(tr["pairs"]), src_words=0, tgt_words=0, num_total=0)
    for i, p in enumerate(tr["pairs"], 1):
        a, b = p[src], p[tgt]
        res["src_words"] += len(a.replace("`", "").split())
        res["tgt_words"] += len(b.replace("`", "").split())
        na, nb = numbers(a, src), numbers(b, tgt)
        res["num_total"] += sum(na.values())
        if na != nb:
            fail(f"{tr['header']['title']}: numbers differ in paragraph {i}: "
                 f"only in source {dict(na - nb)}, only in translation {dict(nb - na)}")
    tgt_text = " ".join(p[tgt] for p in tr["pairs"])
    visible = re.sub(r"`[^`]*`", " ", tgt_text)
    if tgt == "en":
        left = re.findall(r"[А-Яа-яЁё]+", visible)
        res["leftover"] = f"Cyrillic words in the English text: {len(left)}"
    else:
        left = re.findall(r"[A-Za-z]{2,}", visible)
        res["leftover"] = f"Latin words in the Russian text outside the verbatim citation: {len(left)}"
    if left:
        fail(f"{tr['header']['title']}: untranslated words {left[:10]}")
    used = [glossary_used(tgt_text, row[1]) for row in tr["glossary"]]
    res["gl_used"], res["gl_total"] = sum(used), len(used)
    for row, ok in zip(tr["glossary"], used):
        if not ok:
            fail(f"{tr['header']['title']}: glossary term not found in the translation: {row[1]}")
    return res


# --------------------------------------------------------------------------- docx helpers
def set_font(style, name, size=None, bold=None, color=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def new_doc(font="Calibri", size=11, landscape=False, footer="", title=""):
    d = Document()
    set_font(d.styles["Normal"], font, size)
    d.styles["Normal"].paragraph_format.space_after = Pt(6)
    d.styles["Normal"].paragraph_format.line_spacing = 1.15
    set_font(d.styles["Title"], font, 22, True, HEAD)
    set_font(d.styles["Heading 1"], font, 16, True, HEAD)
    set_font(d.styles["Heading 2"], font, 13, True, HEAD)
    set_font(d.styles["List Bullet"], font, size)
    sec = d.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(1.8)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    fp = sec.footer.paragraphs[0]
    r = fp.add_run(footer)
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.title = title
    cp.created = cp.modified = BUILD_DT
    return d


def para(doc, text="", style=None, size=None, italic=False, bold=False, color=None, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.italic, r.bold = italic, bold
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return p


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def table(doc, header, rows, widths_cm, font_size=10):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font_size)
        shade(c, "D9E2F3")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
            r = cells[i].paragraphs[0].add_run(val)
            r.font.size = Pt(font_size)
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
        cant = OxmlElement("w:cantSplit")
        cant.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cant)
    return t


def clean(s):
    return s.replace("`", "")


# --------------------------------------------------------------------------- translation samples
def build_translation(fname, stem, tr, font="Calibri"):
    h = tr["header"]
    src, tgt = tr["src"], tr["tgt"]
    qa = qa_translation(tr)
    label = {"ru": "Русский (оригинал)", "en": "English (source)"}
    tlabel = {"ru": "Русский (перевод)", "en": "English (translation)"}
    direction = "RU → EN" if src == "ru" else "EN → RU"
    foot = FOOTER["ru"] + " | " + FOOTER["en"]

    # --- side by side
    d = new_doc(font, 10.5, landscape=True, footer=foot, title=h["title"])
    para(d, h["title"], "Title")
    para(d, f"Side-by-side translation sample ({direction}) · Параллельный текст", size=11, color=GREY)
    para(d, f"Source / Источник: {h['source_title']}", size=9.5)
    para(d, f"URL: {h['source_url']}", size=9.5)
    para(d, f"Licence / Лицензия: {h['source_licence']}", size=9.5)
    para(d, f"Retrieved / Получено: {h['retrieved']}", size=9.5)
    rows = [(clean(p[src]), clean(p[tgt])) for p in tr["pairs"]]
    table(d, [label[src], tlabel[tgt]], rows, [13.05, 13.05], 10)
    d.add_paragraph()
    para(d, "Глоссарий / Glossary", "Heading 1")
    table(d, tr["gl_head"], tr["glossary"], [6.0, 6.0, 14.1], 9.5)
    para(d, "")
    para(d, "Примечания переводчика / Translator's notes", "Heading 1")
    for n in tr["notes"]:
        para(d, n, "List Bullet", size=10)
    para(d, "Контроль качества / QA pass", "Heading 1")
    for line in [
        f"Numbers: {qa['num_total']} numeric tokens in the source; the multiset of numbers matches the translation in every one of the "
        f"{qa['paras']} paragraphs (automated check, run by build_samples.py; the numerals are compared after normalising "
        f"roman/ordinal forms, thousand separators and decimal signs).",
        f"Untranslated words: {qa['leftover']}.",
        f"Glossary: {qa['gl_used']} of {qa['gl_total']} glossary renderings are present in the translation (consistency check).",
        f"Length: source {qa['src_words']} words, translation {qa['tgt_words']} words.",
        "Names, units and number words: checked by hand against the source (see the notes).",
    ]:
        para(d, line, "List Bullet", size=10)
    p1 = OUT / f"{stem}_side-by-side.docx"
    d.save(p1)

    # --- clean target
    d2 = new_doc(font, 11.5, footer=foot, title=h.get("target_title", h["title"]))
    para(d2, h.get("target_title", h["title"]), "Title")
    for p in tr["pairs"]:
        pp = para(d2, clean(p[tgt]))
        pp.paragraph_format.space_after = Pt(8)
    para(d2, "")
    para(d2, f"Source: {h['source_title']}. {h['source_url']}", size=9, color=GREY)
    para(d2, f"Licence of the source: {h['source_licence']}", size=9, color=GREY)
    para(d2, "Translation: Максим (Maxim), 2026. Portfolio demonstration sample; not client work.", size=9, color=GREY)
    p2 = OUT / f"{stem}_clean-{tgt.upper()}.docx"
    d2.save(p2)
    print(f"  {fname}: {qa['paras']} paragraphs, {qa['src_words']} source words, {qa['tgt_words']} target words, "
          f"{qa['num_total']} numbers matched, glossary {qa['gl_used']}/{qa['gl_total']}, {qa['leftover']}")
    return [p1, p2]


# --------------------------------------------------------------------------- editing samples
STYLE = {"t": "Heading 1", "h1": "Heading 1", "h2": "Heading 2", "li": "List Bullet", "body": None}


def mark_inserted_paragraph(p, ids):
    ppr = p._p.get_or_add_pPr()
    rpr = OxmlElement("w:rPr")
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), str(next(ids)))
    ins.set(qn("w:author"), AUTHOR)
    ins.set(qn("w:date"), BUILD_ISO)
    rpr.append(ins)
    ppr.append(rpr)


def tracked_run(p, text, kind, ids):
    el = OxmlElement("w:ins" if kind == "ins" else "w:del")
    el.set(qn("w:id"), str(next(ids)))
    el.set(qn("w:author"), AUTHOR)
    el.set(qn("w:date"), BUILD_ISO)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t" if kind == "ins" else "w:delText")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    el.append(r)
    p._p.append(el)


def render_edit(doc, paras, mode, ids):
    """mode: draft | clean | markup | tracked.  Returns counts of changes."""
    stat = Counter()
    for kind, tokens in paras:
        if mode == "draft" and not derive(tokens, "draft").strip():
            continue  # a paragraph the editor added (e.g. a new heading) is not in the draft
        pars = [doc.add_paragraph(style=STYLE.get(kind))]
        last = None
        lstrip = False
        for t, x in tokens:
            cur = pars[-1]
            if t == "com":
                if mode == "markup" and last is not None:
                    doc.add_comment(runs=last, text=x, author=AUTHOR, initials=INITIALS)
                    stat["comments"] += 1
                continue
            if t == "ins" and x == "¶":
                stat["splits"] += 1
                if mode == "draft":
                    cur.add_run(" ")
                    continue
                if mode == "markup":
                    r = cur.add_run("¶")
                    r.font.color.rgb = GREEN
                    last = r
                elif mode == "tracked":
                    mark_inserted_paragraph(cur, ids)
                pars.append(doc.add_paragraph())
                lstrip = True
                continue
            if t == "del" and mode == "clean":
                continue
            if t == "ins" and mode == "draft":
                continue
            if lstrip and t in ("keep", "ins"):
                x = x.lstrip()
                if not x:
                    continue
                lstrip = False
            elif lstrip:
                lstrip = False
            if t == "del":
                stat["deletions"] += 1
            elif t == "ins":
                stat["insertions"] += 1
            if mode == "tracked" and t in ("del", "ins"):
                tracked_run(cur, x, t, ids)
                continue
            r = cur.add_run(x)
            if mode == "markup":
                if t == "del":
                    r.font.strike = True
                    r.font.color.rgb = RED
                elif t == "ins":
                    r.font.underline = True
                    r.font.color.rgb = GREEN
            last = r
    return stat


def render_note(doc, note):
    for kind, text in note:
        if kind == "h":
            para(doc, text, "Heading 1")
        elif kind == "h2":
            para(doc, text, "Heading 2")
        elif kind == "li":
            para(doc, text, "List Bullet")
        else:
            para(doc, text)


def build_edit(fname, stem, ed, lang, font):
    h = ed["header"]
    paras = ed["paras"]
    foot = FOOTER[lang]
    title = h["title"]
    drafts = [derive(t, "draft") for _, t in paras]
    edited = [derive(t, "edited") for _, t in paras]
    dw = sum(len(x.replace("\n", " ").split()) for k, x in zip([k for k, _ in paras], drafts) if k != "t")
    ew = sum(len(x.replace("\n", " ").split()) for k, x in zip([k for k, _ in paras], edited) if k != "t")
    for x in drafts + edited:
        if "{" in x or "}" in x or "^[" in x:
            fail(f"{fname}: stray markup left in text: {x[:60]!r}")
    labels = {
        "ru": dict(legend="Условные обозначения: красный зачёркнутый текст — удалено редактором; зелёный подчёркнутый — добавлено; "
                          "«¶» — разбивка абзаца; пояснения — в комментариях на полях.",
                   draft="Авторский черновик (исходный текст)", edited="Отредактированный текст",
                   markup="Правка с пометками", tracked="Правка в режиме исправлений Word",
                   stats="Объём: черновик — {dw} сл., после правки — {ew} сл. · удалений: {deletions}, вставок: {insertions}, "
                         "разбивок абзаца: {splits}, комментариев: {comments}.",
                   summary="Сводка изменений", note_title="Заметка редактора"),
        "en": dict(legend="Legend: red strikethrough = deleted by the editor; green underline = added; "
                          "\"¶\" = paragraph split; explanations are in the margin comments.",
                   draft="Author's draft (as received)", edited="Edited text",
                   markup="Edit with visible markup", tracked="Edit as Word tracked changes",
                   stats="Scope: draft {dw} words, edited {ew} words · deletions: {deletions}, insertions: {insertions}, "
                         "paragraph splits: {splits}, comments: {comments}.",
                   summary="Change summary", note_title="Editor's note"),
    }[lang]
    outs = []

    def cover(doc, sub):
        para(doc, title, "Title")
        para(doc, sub, size=12, color=GREY)
        para(doc, h["about"], size=9.5, italic=True, color=GREY)

    # draft & clean
    for mode, sub, suffix in (("draft", labels["draft"], "draft"), ("clean", labels["edited"], "edited")):
        d = new_doc(font, 12 if lang == "ru" else 11, footer=foot, title=title)
        cover(d, sub)
        render_edit(d, paras, mode, count(1))
        pth = OUT / f"{stem}_{suffix}.docx"
        d.save(pth)
        outs.append(pth)
    # markup & tracked
    statlast = None
    for mode, sub, suffix in (("markup", labels["markup"], "markup"), ("tracked", labels["tracked"], "tracked-changes")):
        d = new_doc(font, 12 if lang == "ru" else 11, footer=foot, title=title)
        cover(d, sub)
        if mode == "markup":
            para(d, labels["legend"], size=9.5, color=GREY)
        stat = render_edit(d, paras, mode, count(1))
        if mode == "markup":
            statlast = stat
            para(d, labels["stats"].format(dw=dw, ew=ew, deletions=stat["deletions"], insertions=stat["insertions"],
                                           splits=stat["splits"], comments=stat["comments"]), size=9.5, color=GREY)
        if ed["summary"]:
            d.add_page_break()
            para(d, labels["summary"], "Heading 1")
            sh, sr = ed["summary"]
            table(d, sh, sr, [3.0, 7.4, 5.2], 9.5)
        if ed["note"]:
            d.add_page_break()
            if not (ed["note"] and ed["note"][0][0] == "h"):
                para(d, labels["note_title"], "Heading 1")
            render_note(d, ed["note"])
        pth = OUT / f"{stem}_{suffix}.docx"
        d.save(pth)
        outs.append(pth)
    print(f"  {fname}: draft {dw} words -> edited {ew} words; {statlast['deletions']} deletions, "
          f"{statlast['insertions']} insertions, {statlast['splits']} paragraph splits, {statlast['comments']} comments")
    return outs


# --------------------------------------------------------------------------- verification
def verify(paths):
    ok = True
    for p in paths:
        try:
            with zipfile.ZipFile(p) as z:
                bad = z.testzip()
                names = z.namelist()
                xml = z.read("word/document.xml").decode("utf-8")
            doc = Document(p)
            extra = ""
            if "tracked" in p.name:
                extra = f", w:ins={xml.count('<w:ins ')} w:del={xml.count('<w:del ')}"
            if "markup" in p.name:
                extra = f", comments part={'word/comments.xml' in names}"
            if bad:
                ok = False
            print(f"  OK {p.name}: {p.stat().st_size} bytes, {len(doc.paragraphs)} paragraphs, "
                  f"{len(doc.tables)} tables{extra}")
        except Exception as e:  # noqa
            ok = False
            print(f"  BROKEN {p.name}: {e}")
    return ok


def main():
    OUT.mkdir(exist_ok=True)
    made = []
    print("Translations")
    made += build_translation("01_ru-en_samovar.md", "01_RU-EN_samovar",
                              parse_translation(SRC / "01_ru-en_samovar.md"))
    made += build_translation("02_en-ru_groundwater.md", "02_EN-RU_groundwater",
                              parse_translation(SRC / "02_en-ru_groundwater.md"))
    print("Editing")
    made += build_edit("03_literary-editing_ru.md", "03_literary-editing_RU",
                       parse_edit(SRC / "03_literary-editing_ru.md"), "ru", "Times New Roman")
    made += build_edit("04_business-editing_en.md", "04_business-editing_EN",
                       parse_edit(SRC / "04_business-editing_en.md"), "en", "Calibri")
    print("Verification")
    ok = verify(made)
    if errors or not ok:
        print(f"\nFAILED: {len(errors)} QA problem(s)")
        sys.exit(1)
    print(f"\nDone: {len(made)} DOCX files written to {OUT}")


if __name__ == "__main__":
    main()
